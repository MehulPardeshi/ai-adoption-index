"""Convert a filled labeling workbook into rater_<name>.csv (for src/eval/irr.py), with checks.

Usage (repo root):  python3 -m src.eval.workbook_to_csv path/to/labeling_<name>.xlsx <name> [round, default 1]
Checks: every item has an evidence code; codes are valid; scenarios required on D2/D1/P and blank on N/X; scenario codes known.
"""
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

LEVELS = {"D2", "D1", "P", "N", "X"}
CODES = {"CX", "ADV", "CREDIT", "RISKM", "FRAUD", "SURV", "CYBER", "DEV", "PROD", "OPS", "MKT", "ENT",
         "GOV", "TALENT", "INVEST"}


def main(path: str, name: str, round_no: int = 1) -> None:
    ws = load_workbook(path, data_only=True)["Label here"]
    rows, problems = [], []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[0] is None:
            continue
        item, pid, about, ev, sc, note = r[0], r[1], (r[7] or ""), (r[8] or ""), (r[9] or ""), (r[10] or "")
        ev = str(ev).strip().upper()
        codes = [c.strip().upper() for c in str(sc).replace(",", ";").split(";") if c.strip()]
        if ev not in LEVELS:
            problems.append(f"item {item}: evidence '{ev or 'blank'}'")
        if codes and ev in ("N", "X"):
            problems.append(f"item {item}: scenarios given but evidence is {ev}")
        if ev in ("D2", "D1", "P") and not codes:
            problems.append(f"item {item}: {ev} needs at least one scenario")
        bad = [c for c in codes if c not in CODES]
        if bad:
            problems.append(f"item {item}: unknown scenario code(s) {bad}")
        rows.append([item, pid, str(about).strip().upper(), ev, ";".join(codes), note])
    if problems:
        print("Fix these first:\n  " + "\n  ".join(problems))
        sys.exit(1)
    out = Path(f"docs/part3-metrics/labeling_round_{round_no}") / f"rater_{name.lower()}.csv"
    with out.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["item", "passage_id", "about_ai", "evidence", "scenarios", "note"])
        w.writerows(rows)
    print(f"OK: {len(rows)} items -> {out}")


if __name__ == "__main__":
    if len(sys.argv) not in (3, 4):
        print(__doc__)
        sys.exit(2)
    main(sys.argv[1], sys.argv[2], int(sys.argv[3]) if len(sys.argv) == 4 else 1)
