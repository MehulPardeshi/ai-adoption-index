"""Build an Excel workbook for a labeling round: dropdowns, progress counter, rule checks.

The workbook contains NO labels: raters stay blind. Each rater saves a copy as
labeling_<name>.xlsx, fills it in, then converts it with
    python3 -m src.eval.workbook_to_csv labeling_<name>.xlsx <name>
which writes docs/part3-metrics/labeling_round_1/rater_<name>.csv for src/eval/irr.py
(add the round number for later rounds: ... labeling_<name>.xlsx <name> 2).

Run from the repo root:  python3 -m src.eval.make_labeling_workbook [round, default 1]
The cheat-sheet tab follows labeling_round_2/cheat_sheet.md (v2, Oct 27); the committed round-1 workbook keeps v1.
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SCENARIOS = "CX, ADV, CREDIT, RISKM, FRAUD, SURV, CYBER, DEV, PROD, OPS, MKT, ENT, GOV, TALENT, INVEST"
HEADER = ["item", "passage_id", "company", "form", "period", "section", "passage",
          "about_ai", "evidence", "scenarios", "note"]
CHEAT = [
    ("EVIDENCE LEVEL: does THIS company say IT uses AI?", ""),
    ("ABOUT AI?", "Y for anything about AI, INCLUDING AI risks, criminals' or competitors' AI, AI laws and trends (those get evidence N). "
                  "N only if the keyword matched something that isn't AI (e.g. 'predictive models' for ordinary scorecards) -> evidence X."),
    ("", ""),
    ("D2", "Deployed AND specific: names a tool, use case, number, oversight body, role or date. e.g. 'Our AI assistant answered 3 million customer questions this year.'"),
    ("D1", "Deployed but general: says it uses or governs AI, no specifics. e.g. 'We use AI in parts of our business.'"),
    ("P", "Planned or expected use. e.g. 'We intend to expand our use of AI next year.'"),
    ("N", "No evidence of own use: risks, threats, competitors, laws, trends, hedged 'MAY use'. e.g. 'Criminals may use AI to commit fraud.'"),
    ("X", "Not substantive: definition, acronym list, heading, or not actually about AI (about_ai = N). e.g. 'AI: artificial intelligence.'"),
    ("", ""),
    ("TIE-BREAKERS", "Mixed passage -> highest level. A risk-factor passage is still D1 if it says the company itself uses AI. "
                     "'may use / could include' -> N. Other companies, criminals, the economy -> N. "
                     "Forward-looking-statement lists ('...technology changes, including AI') -> N. "
                     "Heading alone -> X, even if it says 'our use of AI'. about_ai = N -> X."),
    ("", ""),
    ("SCENARIOS (required for D2, D1 and P; blank for N/X; separate with ;)", SCENARIOS),
    ("CX", "customer service"), ("ADV", "advice / wealth / retirement"), ("CREDIT", "lending / underwriting"),
    ("RISKM", "risk & finance models"), ("FRAUD", "fraud / AML / KYC"), ("SURV", "compliance & trade surveillance"),
    ("CYBER", "cyber defense"), ("DEV", "software engineering"), ("PROD", "employee productivity tools"),
    ("OPS", "operations / data / reporting"), ("MKT", "markets & trading"), ("ENT", "enterprise-wide, area not named"),
    ("GOV / TALENT / INVEST", "enablers: governance & leadership / training & talent / investment & M&A"),
    ("Rule (frozen Oct 27)", "A general sentence listing several areas ('we use AI for A, B and C') counts ONCE, as ENT."),
    ("", ""),
    ("RULES OF THE ROUND", "Label on your own, with your own judgment (no AI suggestions). "
                           "Don't open passage_labels_v0.csv, company_scores_v0.csv or agent_labels_*.csv. "
                           "Don't compare answers until all three raters are done. If unsure, pick your best answer and explain in 'note'."),
]


def main(round_no: int = 1) -> None:
    ROUND = Path(f"docs/part3-metrics/labeling_round_{round_no}")
    sample = list(csv.DictReader((ROUND / "sample_passages.csv").open(encoding="utf-8")))
    wb = Workbook()
    ws = wb.active
    ws.title = "Label here"
    ws.append([f"Labeling round {round_no}: fill the 4 yellow columns. Progress:", "", "", "", "", "",
               f'=COUNTA(I3:I{len(sample) + 2})&" / {len(sample)} labeled"'])
    ws["A1"].font = Font(bold=True, size=12)
    ws["G1"].font = Font(bold=True, size=12, color="1F4E78")
    ws.append(HEADER)
    head_fill, todo_fill = PatternFill("solid", fgColor="1F4E78"), PatternFill("solid", fgColor="FFF2CC")
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    for row in sample:
        ws.append([int(row["item"]), row["passage_id"], row["company"], row["form"], row["period"],
                   row["section_guess"], row["passage"], "", "", "", ""])
    last = len(sample) + 2
    widths = {"A": 6, "B": 14, "C": 16, "D": 7, "E": 11, "F": 18, "G": 80, "H": 10, "I": 10, "J": 22, "K": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["B"].hidden = True  # passage_id kept for the converter
    for r in range(3, last + 1):
        for col in "ABCDEFGHIJK":
            ws[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        for col in "HIJK":
            ws[f"{col}{r}"].fill = todo_fill
        words = len(str(ws[f"G{r}"].value).split())
        ws.row_dimensions[r].height = max(30, min(400, 15 * (words // 14 + 1)))
    dv_ai = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showDropDown=False)
    dv_ev = DataValidation(type="list", formula1='"D2,D1,P,N,X"', allow_blank=True, showDropDown=False)
    dv_ev.prompt, dv_ev.promptTitle = "D2 specific use · D1 general use · P plan · N no own use · X not substantive", "Evidence level"
    dv_sc = DataValidation(type="custom", formula1="TRUE", allow_blank=True)
    dv_sc.prompt, dv_sc.promptTitle = "Required for D2/D1/P, blank for N/X. Codes separated by ; e.g. CX;FRAUD", "Scenarios"
    for dv, col in ((dv_ai, "H"), (dv_ev, "I"), (dv_sc, "J")):
        ws.add_data_validation(dv)
        dv.add(f"{col}3:{col}{last}")
    red = PatternFill("solid", fgColor="F8CBAD")
    # scenarios filled although evidence is N or X -> red
    ws.conditional_formatting.add(f"J3:J{last}", FormulaRule(
        formula=[f'AND(LEN($J3)>0,OR($I3="N",$I3="X"))'], fill=red))
    # D2/D1/P without scenarios -> red
    ws.conditional_formatting.add(f"J3:J{last}", FormulaRule(
        formula=[f'AND(OR($I3="D2",$I3="D1",$I3="P"),LEN($J3)=0)'], fill=red))
    ws.freeze_panes = "A3"  # keep the header visible; answer columns stay reachable on small screens

    cs = wb.create_sheet("Cheat sheet")
    cs.column_dimensions["A"].width = 34
    cs.column_dimensions["B"].width = 120
    for a, b in CHEAT:
        cs.append([a, b])
        cs[f"A{cs.max_row}"].font = Font(bold=True)
        cs[f"B{cs.max_row}"].alignment = Alignment(wrap_text=True, vertical="top")
    out = ROUND / "labeling_workbook.xlsx"
    wb.save(out)
    print(f"Wrote {out} ({len(sample)} passages, no labels).")


if __name__ == "__main__":
    main(int(sys.argv[1]) if len(sys.argv) > 1 else 1)
